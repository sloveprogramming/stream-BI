import uuid
from pathlib import Path

import pytest
import sqlalchemy as sa
from apscheduler.schedulers.background import BackgroundScheduler
from fastapi import Request
from fastapi.testclient import TestClient

from app.api.deps import Context, get_ctx
from app.bootstrap import bootstrap_admin
from app.main import create_app
from app.models import (
    ApiKey,
    AuditLog,
    Dataset,
    DatasetGrant,
    ReportRun,
    ReportTask,
    Role,
    User,
)
from app.reports.scheduler import refresh_report_jobs
from app.reports.service import generate_report_file


class FakeSettings:
    rate_limit_per_minute = 1000
    rate_limit_per_day = 100000
    export_max_rows = 50000


@pytest.fixture
def env(pg_session_factory, pg_engine):
    suffix = uuid.uuid4().hex[:8]
    _, admin_key = bootstrap_admin(f"repadm_{suffix}", "pw123456")
    admin_headers = {"X-API-Key": admin_key}
    app = create_app()

    def override_ctx(request: Request):
        return Context(
            session_factory=pg_session_factory,
            erp_engine=object(),
            pg_engine=pg_engine,
            settings=FakeSettings(),
        )

    from app.api.deps import get_ctx as _g

    app.dependency_overrides[get_ctx] = override_ctx
    client = TestClient(app)

    with pg_engine.begin() as conn:
        conn.execute(sa.text("CREATE SCHEMA IF NOT EXISTS ods"))
        conn.execute(sa.text(
            "CREATE TABLE IF NOT EXISTS ods.rep_t "
            "(id integer PRIMARY KEY, region text, amount numeric, cost numeric)"
        ))
        conn.execute(sa.text(
            "INSERT INTO ods.rep_t VALUES (1,'华东',10,5),(2,'华北',20,8) "
            "ON CONFLICT (id) DO NOTHING"
        ))
    r = client.post("/api/v1/admin/roles", json={"name": f"r_{suffix}"}, headers=admin_headers)
    role_id = r.json()["id"]
    r = client.post("/api/v1/admin/datasets", json={
        "name": f"repd_{suffix}", "source_table": "ods.rep_t",
        "pk_columns": ["id"], "columns": ["id", "region", "amount", "cost"],
    }, headers=admin_headers)
    ds_id, ds_name = r.json()["id"], r.json()["name"]
    client.post("/api/v1/admin/grants", json={
        "role_id": role_id, "dataset_id": ds_id,
        "row_rules": [{"field": "region", "op": "eq", "value": "华东"}],
        "field_whitelist": ["id", "region", "amount"],
    }, headers=admin_headers)
    r = client.post("/api/v1/admin/users", json={
        "username": f"u_{suffix}", "password": "pw123456", "role_id": role_id,
    }, headers=admin_headers)
    user_id = r.json()["id"]

    def cleanup():
        with pg_session_factory() as s:
            s.query(ReportRun).filter(
                ReportRun.task_id.in_(
                    s.query(ReportTask.id).filter_by(owner_user_id=user_id)
                )
            ).delete(synchronize_session=False)
            s.query(ReportTask).filter_by(owner_user_id=user_id).delete()
            s.query(AuditLog).filter_by(user_id=user_id).delete()
            s.query(ApiKey).filter_by(user_id=user_id).delete()
            s.query(User).filter_by(id=user_id).delete()
            s.query(DatasetGrant).filter_by(role_id=role_id).delete()
            s.query(Role).filter_by(id=role_id).delete()
            s.query(Dataset).filter_by(id=ds_id).delete()
            adm = s.query(User).filter_by(username=f"repadm_{suffix}").first()
            if adm is not None:
                s.query(ApiKey).filter_by(user_id=adm.id).delete()
                s.delete(adm)
            s.commit()

    return {
        "client": client, "admin_headers": admin_headers, "user_id": user_id,
        "role_id": role_id, "ds_id": ds_id, "ds_name": ds_name,
        "suffix": suffix, "cleanup": cleanup, "pg_session_factory": pg_session_factory,
    }


def test_service_generates_scoped_file(env, pg_engine):
    with env["pg_session_factory"]() as s:
        task = ReportTask(
            name=f"svc_{env['suffix']}", dataset_id=env["ds_id"], format="xlsx",
            cron="0 8 * * *", filters={}, owner_user_id=env["user_id"],
        )
        s.add(task)
        s.commit()
        run = generate_report_file(task)
        assert run.status == "success"
        assert run.row_count == 1  # 只有华东
        path = Path("exports") / run.file_name
        assert path.exists()
        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0] == ("id", "region", "amount")
        assert rows[1][1] == "华东"
        path.unlink()
        s.delete(run)
        s.delete(task)
        s.commit()


def test_report_task_api_and_download(env):
    client, headers = env["client"], env["admin_headers"]
    # 用户登录拿 JWT
    r = client.post("/api/v1/auth/login", json={
        "username": f"u_{env['suffix']}", "password": "pw123456",
    })
    token = r.json()["token"]
    web_headers = {"Authorization": f"Bearer {token}"}
    # 用户本人创建报表任务（JWT）
    r = client.post("/api/v1/my/report-tasks", json={
        "name": f"api_{env['suffix']}", "dataset_id": env["ds_id"],
        "format": "csv", "cron": "0 8 * * *",
    }, headers=web_headers)
    assert r.status_code == 201, r.text
    task_id = r.json()["id"]
    # 我的报表（当前为空）
    r = client.get("/api/v1/my/reports", headers=web_headers)
    assert r.status_code == 200 and r.json() == []
    # 手动触发（走任务执行入口）
    from app.reports.service import run_report_task

    with env["pg_session_factory"]() as s:
        task = s.get(ReportTask, task_id)
        run = run_report_task(task)
        assert run.status == "success"
        run_id = run.id
        file_name = run.file_name
    # 我的报表出现
    r = client.get("/api/v1/my/reports", headers=web_headers)
    assert any(item["id"] == run_id for item in r.json())
    # 下载自己的报表
    r = client.get(f"/api/v1/my/reports/{run_id}/download", headers=web_headers)
    assert r.status_code == 200
    # 管理员（非 owner）不能下载
    r = client.get(f"/api/v1/my/reports/{run_id}/download", headers=headers)
    assert r.status_code == 403
    Path("exports", file_name).unlink()


def test_refresh_report_jobs_registers_and_removes(env, pg_session_factory):
    scheduler = BackgroundScheduler(timezone="Asia/Hong_Kong")
    scheduler.start()
    try:
        client, headers = env["client"], env["admin_headers"]
        r = client.post("/api/v1/my/report-tasks", json={
            "name": f"sch_{env['suffix']}", "dataset_id": env["ds_id"],
            "format": "csv", "cron": "0 8 * * *",
        }, headers={"Authorization": f"Bearer {client.post('/api/v1/auth/login', json={'username': f'u_{env['suffix']}', 'password': 'pw123456'}).json()['token']}"})
        task_id = r.json()["id"]
        with pg_session_factory() as s:
            refresh_report_jobs(scheduler, pg_session_factory)
            ids = {j.id for j in scheduler.get_jobs()}
            assert f"report_{task_id}" in ids
            s.query(ReportTask).filter_by(id=task_id).delete()
            s.commit()
            refresh_report_jobs(scheduler, pg_session_factory)
            ids = {j.id for j in scheduler.get_jobs()}
            assert f"report_{task_id}" not in ids
    finally:
        scheduler.shutdown(wait=False)


